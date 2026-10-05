# -*- coding: utf-8 -*-
"""
Build the IAD Lab 8 report matching the exact house style and Task.md specifications.
"""
import os
import re
import json
import pandas as pd
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(BASE, "figures")
OUT_DIR = os.path.join(BASE, "report")
OUT_DOCX = os.path.join(OUT_DIR, "ІАД_КН-2327Б_Багрій-Белз_ЛР-8.docx")

os.makedirs(OUT_DIR, exist_ok=True)

BODY_FONT = "Times New Roman"
BODY_SIZE = Pt(14)
LINE_SPACING = 1.15
FIG_WIDTH = Cm(15.0)

# Завантажуємо результати правил для таблиці
with open(os.path.join(BASE, "report.json"), "r", encoding="utf-8") as f:
    top20_rules = json.load(f)

top5 = top20_rules[:5]

def _set_font(run, family, size, bold=False):
    run.font.name = family
    run.font.size = size
    run.bold = bold
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), family)
    rFonts.set(qn("w:hAnsi"), family)
    rFonts.set(qn("w:cs"), family)

def _para(doc, text="", *, align=None, font=BODY_FONT, size=BODY_SIZE,
          bold=False, line_spacing=LINE_SPACING, keep_with_next=False, keep_together=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    if align is not None:
        p.alignment = align
    if keep_with_next:
        p.paragraph_format.keep_with_next = True
    if keep_together:
        p.paragraph_format.keep_together = True
    if text:
        run = p.add_run(text)
        _set_font(run, font, size, bold=bold)
    return p

def _add_bold_spans(p, text, font=BODY_FONT, size=BODY_SIZE):
    """Parses **bold** markdown tags."""
    pos = 0
    for m in re.finditer(r"\*\*([^*]+)\*\*", text):
        if m.start() > pos:
            run = p.add_run(text[pos:m.start()])
            _set_font(run, font, size, bold=False)
        run = p.add_run(m.group(1))
        _set_font(run, font, size, bold=True)
        pos = m.end()
    if pos < len(text):
        run = p.add_run(text[pos:])
        _set_font(run, font, size, bold=False)

def _add_rich_para(doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = LINE_SPACING
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.alignment = align
    _add_bold_spans(p, text)
    return p

def _add_image(doc, path, width=FIG_WIDTH):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    if os.path.exists(path):
        p.add_run().add_picture(path, width=width)
    else:
        print(f"Warning: image {path} not found")

def _set_cell_border(cell, **kwargs):
    """
    kwargs: top, bottom, left, right
    values: dict(sz=12, val='single', color='FF0000', space='0')
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key in ["sz", "val", "color", "space", "shadow"]:
                if key in edge_data:
                    element.set(qn('w:{}'.format(key)), str(edge_data[key]))

def _set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def build_report():
    doc = Document()

    # 1. Поля 1.5 см
    for s in doc.sections:
        s.top_margin = Cm(1.5)
        s.bottom_margin = Cm(1.5)
        s.left_margin = Cm(1.5)
        s.right_margin = Cm(1.5)

    # 2. Титульний блок
    c = WD_ALIGN_PARAGRAPH.CENTER
    _para(doc, "ЗВІТ", align=c)
    _para(doc, "про виконання лабораторної роботи №8", align=c)
    _para(doc, "«Асоціативні правила. Алгоритм Apriori»", align=c)
    _para(doc, "з дисципліни", align=c)
    _para(doc, "«Інтелектуальний аналіз даних»", align=c)
    _para(doc, "Студента групи КН-2327Б", align=c)
    _para(doc, "Багрія-Белза Андрія", align=c, bold=True)

    # 3. Мета роботи
    _add_rich_para(doc, "**Мета роботи:** Дослідити алгоритм видобування асоціативних правил Apriori для аналізу купівельного кошика (Market Basket Analysis), встановити закономірності спільних покупок за метриками підтримки, достовірності та підйому й візуалізувати виявлені закономірності.")

    # 4. Завдання 1
    _add_rich_para(doc, "**Завдання 1 (Алгоритм Apriori):** Проведено попереднє очищення транзакцій датасету Basket_data.csv (7501 чек, 119 найменувань) від порожніх значень NaN. Застосовано алгоритм Apriori з параметрами фільтрації: min_support = 0.003 (поява щонайменше у 23 транзакціях), min_confidence = 0.2 (мінімальна достовірність 20%), min_lift = 3.0 (перевага над випадковою сумісністю у 3+ рази) та max_length = 3. Сформовано 72 валідні асоціативні правила, відсортовані за спаданням показника lift.")
    # Таблиця топ-5 правил
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    hdr_titles = ["Асоціативне правило (Base -> Add)", "Support", "Confidence", "Lift"]
    col_widths = [Cm(7.5), Cm(2.5), Cm(2.8), Cm(2.5)]

    for idx, heading in enumerate(hdr_titles):
        hdr_cells[idx].width = col_widths[idx]
        p = hdr_cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(heading)
        _set_font(run, BODY_FONT, Pt(11), bold=True)
        _set_cell_margins(hdr_cells[idx], top=80, bottom=80, left=100, right=100)
        _set_cell_border(hdr_cells[idx],
                         top=dict(val='single', sz=6, color='999999'),
                         bottom=dict(val='single', sz=12, color='333333'),
                         left=dict(val='none'), right=dict(val='none'))

    for r_idx, rule in enumerate(top5):
        row_cells = table.add_row().cells
        rule_text = f"{rule['base']} -> {rule['add']}"
        sup_text = f"{rule['support']:.4f} ({rule['support']*100:.2f}%)"
        conf_text = f"{rule['confidence']:.3f} ({rule['confidence']*100:.1f}%)"
        lift_text = f"{rule['lift']:.2f}"

        vals = [rule_text, sup_text, conf_text, lift_text]
        aligns = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER]

        for c_idx, val in enumerate(vals):
            row_cells[c_idx].width = col_widths[c_idx]
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = aligns[c_idx]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(val)
            _set_font(run, BODY_FONT, Pt(10.5), bold=(c_idx == 3))
            _set_cell_margins(row_cells[c_idx], top=60, bottom=60, left=100, right=100)
            _set_cell_border(row_cells[c_idx],
                             bottom=dict(val='single', sz=4, color='CCCCCC'),
                             left=dict(val='none'), right=dict(val='none'))

    _para(doc, "") # пустий рядок-відступ після таблиці

    # 5. Завдання 2
    _add_rich_para(doc, "**Завдання 2 (Візуалізація структури правил):** Побудовано орієнтований граф асоціацій (Top-20 за показником Lift), кругову діаграму часток підтримки провідних правил та зіставлення показників Support і Lift.")

    _add_image(doc, os.path.join(FIG, "apriori_rules_graph.png"), width=Cm(15.0))
    _add_rich_para(doc, "**Інтерпретація графа:** Орієнтовані дуги відображають спрямованість впливу: наявність комбінації томатного соусу та спагеті або трав з перцем практично уп'ятеро збільшує шанс купівлі яловичого фаршу (Lift 4.98 та 4.70), а покупка цільнозернової пасти разом із водою стимулює придбання оливкової олії (Lift 6.12).")

    _add_image(doc, os.path.join(FIG, "apriori_pie_chart.png"), width=Cm(13.5))
    _add_rich_para(doc, "**Інтерпретація розподілу підтримки:** Провідні 7 виявлених пар/трійок акумулюють понад 15% від загальної підтримки всіх знайдених специфічних правил, демонструючи стійкі базові кулінарні та кошикові комбінації покупців.")

    _add_image(doc, os.path.join(FIG, "apriori_support_bar.png"), width=Cm(15.0))
    _add_rich_para(doc, "**Бізнес-рекомендації мерчандайзингу:** Виявлені сильні парні зв'язки (паста + ескалоп, легкі вершки + курка, нежирний сир + мед) рекомендується використовувати для крос-викладки супутніх товарів у торгових залах супермаркету, налаштування систем рекомендацій в онлайн-магазині та розробки пакетних промо-акцій.")

    # 6. Висновок
    _add_rich_para(doc, "**Висновок:** Під час виконання лабораторної роботи я опанував алгоритм Apriori для аналізу асоціативних правил купівельного кошика та реалізував повний пайплайн передобробки транзакцій. Дослідив вплив метрик Support, Confidence і Lift на якість видобутих зв'язків та переконався, що орієнтовані графи й зіставлення оцінок дозволяють ефективно виявляти приховані комерційні закономірності для мерчандайзингу та крос-продажів.")

    doc.save(OUT_DOCX)
    print(f"Згенеровано звіт: {OUT_DOCX}")

if __name__ == "__main__":
    build_report()
