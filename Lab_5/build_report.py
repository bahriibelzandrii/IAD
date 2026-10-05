# -*- coding: utf-8 -*-
"""
Build the IAD Lab 5 report (.docx) matching the canonical house style.
"""
import os
import re
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BASE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(BASE, "figures")
OUT_DIR = os.path.join(BASE, "report")
OUT_DOCX = os.path.join(OUT_DIR, "ІАД_КН-2327Б_Багрій-Белз_ЛР-5.docx")

BODY_FONT = "Times New Roman"
BODY_SIZE = Pt(14)
FIG_WIDTH = Cm(14.5)
MARGIN_CM = 1.5


def f(name):
    return os.path.join(FIG, name)


META = {
    "course_name": "Інтелектуальний аналіз даних",
    "course_code": "ІАД",
    "lab_number": 5,
    "topic": "Кластерний аналіз даних (K-Means, ієрархічна кластеризація)",
    "author": "Багрій-Белз Андрій",
    "group": "КН-2327Б",
    "objective": (
        "Опанувати практичні навички застосування методів кластерного аналізу даних у мові Python: "
        "алгоритму K-середніх (K-Means), визначення оптимальної кількості кластерів методами ліктя й силуету, "
        "а також агломеративної ієрархічної кластеризації часових рядів та багатовимірних фінансових показників."
    )
}

BODY = [
    {
        "type": "task",
        "number": 1,
        "text": "Кластеризація методом K-середніх та обґрунтування вибору кількості кластерів"
    },
    {
        "type": "text",
        "text": (
            "Побудовано двовимірний набір даних характеристик фруктів (вага у грамах та ціна в USD/кг). "
            "Навчено модель KMeans із параметрами **n_clusters = 3** та random_state = 42. "
            "Отримано стійкий розподіл спостережень за трьома кластерами та визначено точні координати центроїдів. "
            "Для математичного обґрунтування вибору кількості кластерів k проведено обчислення значень "
            "внутрішньокластерної інерції (WCSS) та середнього коефіцієнта силуету (Silhouette Score) для діапазону k від 1 до 10. "
            "Точка вираженого вигину кривої WCSS («лікоть») та глобальний максимум коефіцієнта силуету (**0.7264**) "
            "однозначно підтверджують оптимальність вибору саме **k = 3** кластерів."
        )
    },
    {
        "type": "result",
        "images": [f("kmeans_clusters.png"), f("elbow_method.png")]
    },
    {
        "type": "task",
        "number": 2,
        "text": "Ієрархічна агломеративна кластеризація часових рядів котирувань акцій (stocks.csv)"
    },
    {
        "type": "text",
        "text": (
            "Виконано ієрархічну агломеративну кластеризацію перших 50 часових спостережень файлу stocks.csv за методом Уорда "
            "(linkage за евклідовою метрикою). Побудована дендрограма відображає хронологічні періоди ринкової динаміки "
            "та високу спорідненість станів ринку в суміжні торговельні дні.\n"
            "Для дослідження взаємозв'язку між самими компаніями проведено попередню стандартизацію часових рядів котирувань "
            "(StandardScaler) та ієрархічну кластеризацію транспонованої матриці фінансових активів. "
            "Отримана дендрограма тікерів об'єднує акції у змістовні секторальні кластери зі схожим профілем ринкової поведінки "
            "(зокрема, спільні групи енергетичних, банківських та металургійних активів)."
        )
    },
    {
        "type": "result",
        "images": [f("stocks_dendrogram.png"), f("tickers_dendrogram.png")]
    }
]

CONCLUSION = (
    "На цій лабораторній роботі я практично дослідив роботу методів неконтрольованого навчання у Python: "
    "алгоритму K-середніх та ієрархічної агломеративної кластеризації. "
    "Експериментально підтвердив ефективність узгодженого використання методу ліктя та метрики силуету для точного "
    "обґрунтування вибору k = 3 (Silhouette Score = 0.7264). "
    "Застосування методу Уорда до стандартизованих фінансових часових рядів дозволило автоматично згрупувати акції за спільними ринковими тенденціями."
)

REPORT = {"meta": META, "body": BODY, "conclusion": CONCLUSION}


def _set_font(run, family, size, bold=False):
    run.font.name = family
    run.font.size = size
    run.bold = bold
    rPr = run._r.get_or_add_rPr()
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), family)
    rFonts.set(qn("w:hAnsi"), family)
    rFonts.set(qn("w:eastAsia"), family)
    rFonts.set(qn("w:cs"), family)


def _para(doc, text="", *, align=None, font=BODY_FONT, size=BODY_SIZE,
          bold=False, line_spacing=1.15, keep_with_next=False, keep_together=True):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if align is not None:
        pf.alignment = align
    pf.line_spacing = line_spacing
    pf.space_before = Pt(0)
    pf.space_after = Pt(4)
    pf.keep_with_next = keep_with_next
    pf.keep_together = keep_together
    if text:
        run = p.add_run(text)
        _set_font(run, font, size, bold=bold)
    return p


def _add_bold(p, text, font=BODY_FONT, size=BODY_SIZE):
    parts = text.split("**")
    for i, seg in enumerate(parts):
        if not seg:
            continue
        if i % 2 == 1:
            _set_font(p.add_run(seg), font, size, bold=True)
        else:
            _set_font(p.add_run(seg), font, size, bold=False)


def _add_image(doc, path):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_together = True
    if os.path.exists(path):
        run = p.add_run()
        run.add_picture(path, width=FIG_WIDTH)
    else:
        print(f"  warning: image not found: {path}")


def _genitive(name):
    parts = name.split()
    if len(parts) < 2:
        return name
    sur, first = parts[0], parts[1]
    sur_sub = []
    for sub in sur.split("-"):
        if sub.endswith(("й", "ї")):
            sur_sub.append(sub[:-1] + "я")
        elif not sub.endswith(("а", "я")):
            sur_sub.append(sub + "а")
        else:
            sur_sub.append(sub)
    sur_gen = "-".join(sur_sub)

    if first.endswith(("й", "ї")):
        first_gen = first[:-1] + "я"
    elif not first.endswith(("а", "я")):
        first_gen = first + "а"
    else:
        first_gen = first
    return f"{sur_gen} {first_gen}"


def _title_block(doc, meta):
    c = WD_ALIGN_PARAGRAPH.CENTER
    lines = [
        "ЗВІТ",
        f"про виконання лабораторної роботи №{meta['lab_number']}",
        f"«{meta['topic']}»",
        "з дисципліни",
        f"«{meta['course_name']}»",
        f"Студента групи {meta['group']}",
        _genitive(meta["author"]),
    ]
    for i, line in enumerate(lines):
        _para(doc, line, align=c, bold=(i == len(lines) - 1))


def build_docx(data, out_docx):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(MARGIN_CM)
    sec.top_margin = sec.bottom_margin = Cm(MARGIN_CM)
    sec.header_distance = sec.footer_distance = Cm(1.25)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(0)

    _title_block(doc, data["meta"])

    if data["meta"].get("objective"):
        p = _para(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY, keep_with_next=True)
        _set_font(p.add_run("Мета роботи: "), BODY_FONT, BODY_SIZE, bold=True)
        _add_bold(p, data["meta"]["objective"])

    J = WD_ALIGN_PARAGRAPH.JUSTIFY
    for block in data["body"]:
        btype = block["type"]
        if btype == "task":
            p = _para(doc, align=J, keep_with_next=True)
            _set_font(p.add_run(f"Завдання {block['number']}: "), BODY_FONT, BODY_SIZE, bold=True)
            _add_bold(p, block.get("text", ""))
        elif btype == "text":
            for line in block["text"].split("\n"):
                line = line.strip()
                if not line:
                    continue
                p = _para(doc, align=J, keep_with_next=True)
                _add_bold(p, line)
        elif btype == "result":
            for img in block["images"]:
                _add_image(doc, img)

    p = _para(doc, align=J)
    _set_font(p.add_run("Висновок: "), BODY_FONT, BODY_SIZE, bold=True)
    _add_bold(p, data["conclusion"])

    Path(out_docx).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_docx)
    print(f"Збережено звіт: {out_docx}")


if __name__ == "__main__":
    build_docx(REPORT, OUT_DOCX)
