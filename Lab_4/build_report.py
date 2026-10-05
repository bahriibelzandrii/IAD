# -*- coding: utf-8 -*-
"""Генератор офіційного звіту DOCX для Лабораторної роботи №4 (ІАД).

Канонічний стиль Zvit (report-maker):
- Times New Roman 14 pt, інтервал 1.15, поля 1.5 см.
- Титульний блок на першій сторінці (7 рядків, центровано).
- Лаконічний текст без води.
- Графіки високої роздільної здатності (ширина 14.5-15.0 см, центровані).
- 1 абзац висновку у минулому часі від першої особи.
"""

import os
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BASE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(BASE, "figures")
OUT_DIR = os.path.join(BASE, "report")
OUT_DOCX = os.path.join(OUT_DIR, "ІАД_КН-2327Б_Багрій-Белз_ЛР-4.docx")

BODY_FONT = "Times New Roman"
BODY_SIZE = Pt(14)
FIG_WIDTH = Cm(15.0)
MARGIN_CM = 1.5


def f(name):
    return os.path.join(FIG, name)


META = {
    "doc_type": "ЗВІТ",
    "sub_title": "про виконання лабораторної роботи №4",
    "topic": "Методи класифікації. Дерева рішень у Python",
    "course_label": "з дисципліни",
    "course": "«Інтелектуальний аналіз даних»",
    "student_group": "Студента групи КН-2327Б",
    "student_name": "Багрія-Белза Андрія",
}

BODY = [
    {
        "type": "text",
        "text": "**Мета роботи:** Ознайомитися з алгоритмами побудови дерев рішень (DecisionTreeClassifier) для задач бінарної класифікації у бібліотеці scikit-learn, дослідити критерій ентропії, методи регуляризації та обрізки (pruning) дерева через оптимізацію гіперпараметрів за допомогою 5-кратної крос-валідації.",
    },
    {
        "type": "text",
        "text": "**Завдання 1. Попередній аналіз вибірки та побудова базового дерева рішень:** Завантажено набір даних медичної діагностики серцево-судинних захворювань heart.xlsx (303 зразки, 5 факторів: sex, cp, fbs, restecg, exang; цільовий клас target: 138 здорових, 165 хворих). Вибірку розбито на навчальну (70%, 212 спостережень) та тестову (30%, 91 спостереження) із фіксованим random_state=42 та стратифікацією. Навчано повне дерево з критерієм ентропії (criterion='entropy'). Отримана модель має глибину 5 рівнів та 22 кінцеві листки, показуючи точність Accuracy = 0.7033 (70.33%), Precision = 0.7091, Recall = 0.7800, F1 = 0.7429.",
    },
    {
        "type": "image",
        "path": f("tree_full.png"),
    },
    {
        "type": "text",
        "text": "**Завдання 2. Оптимізація та регуляризація дерева через крос-валідацію:** Повне дерево схильне до перенавчання (overfitting) на локальних шумових ознаках навчальної вибірки. Для підбору мінімальної кількості зразків для розщеплення вузла (min_samples_split) проведено 5-кратну крос-валідацію (cv=5) у діапазоні від 2 до 50 зразків. Оптимальне значення, що забезпечує мінімальну помилку класифікації (CV Error = 0.2166, середня точність CV Accuracy = 78.34%), становить min_samples_split = 37.",
    },
    {
        "type": "image",
        "path": f("cv_error_plot.png"),
    },
    {
        "type": "text",
        "text": "**Завдання 3. Побудова обрізаного дерева та оцінка узагальнюючої здатності:** Навчано редуковану модель з параметром min_samples_split = 37. Кількість листків зменшилася вдвічі (з 22 до 11), що суттєво спростило інтерпретованість правил рішень без втрати загальної точності: Accuracy = 0.7033, Recall зріс до 0.8200 (виявлено 41 випадок захворювання з 50), а показник F1-score зріс з 0.7429 до 0.7523. Теплові карти матриць невідповідностей демонструють зменшення хибнонегативних помилок (False Negatives скоротилися з 11 до 9).",
    },
    {
        "type": "image",
        "path": f("tree_pruned.png"),
    },
    {
        "type": "image",
        "path": f("confusion_matrix.png"),
    },
]

CONCLUSION = (
    "Висновок: Під час виконання лабораторної роботи я опанував практичні навички "
    "побудови та візуалізації класифікаційних дерев рішень на базі scikit-learn з використанням "
    "критерію приросту інформації (ентропії). На прикладі медичного датасету heart.xlsx я дослідив "
    "проблему перенавчання нерегуляризованого дерева та реалізував процедуру обрізки через підбір "
    "гіперпараметра min_samples_split за допомогою 5-кратної крос-валідації. В результаті оптимізації "
    "вдалося спростити архітектуру дерева вдвічі (з 22 до 11 листків), підвищивши повноту виявлення хворих "
    "(Recall з 0.78 до 0.82) та інтегральну метрику F1-score до 0.7523 при збереженні базової точності."
)

REPORT = {"meta": META, "body": BODY, "conclusion": CONCLUSION}


def _set_font(run, family, size, bold=False):
    run.font.name = family
    run.font.size = size
    run.font.bold = bold
    rPr = run._r.get_or_add_rPr()
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), family)
    rFonts.set(qn("w:hAnsi"), family)
    rFonts.set(qn("w:cs"), family)
    rPr.append(rFonts)


def _para(
    doc,
    text="",
    *,
    align=None,
    font=BODY_FONT,
    size=BODY_SIZE,
    bold=False,
    line_spacing=1.15,
    keep_with_next=False,
    keep_together=True,
):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = keep_with_next
    p.paragraph_format.keep_together = keep_together
    if align is not None:
        p.alignment = align
    if text:
        run = p.add_run(text)
        _set_font(run, font, size, bold=bold)
    return p


def _add_bold(p, text, font=BODY_FONT, size=BODY_SIZE):
    parts = text.split("**")
    for i, seg in enumerate(parts):
        if not seg:
            continue
        if i % 2 == 1:
            _set_font(p.add_run(seg), font, size, bold=True)
        else:
            _set_font(p.add_run(seg), font, size, bold=False)


def _add_image(doc, path):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = False
    p.paragraph_format.keep_together = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists(path):
        p.add_run().add_picture(path, width=FIG_WIDTH)
    else:
        print(f"  warning: image not found: {path}")


def _title_block(doc, meta):
    c = WD_ALIGN_PARAGRAPH.CENTER
    lines = [
        meta["doc_type"],
        meta["sub_title"],
        f"«{meta['topic']}»",
        meta["course_label"],
        meta["course"],
        meta["student_group"],
        meta["student_name"],
    ]
    for i, line in enumerate(lines):
        _para(doc, line, align=c, bold=(i == len(lines) - 1))


def build_docx(data, out_docx):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(MARGIN_CM)
        section.bottom_margin = Cm(MARGIN_CM)
        section.left_margin = Cm(MARGIN_CM)
        section.right_margin = Cm(MARGIN_CM)
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)

    # 1. Титульний блок
    _title_block(doc, data["meta"])

    # 2. Тіло звіту
    for item in data["body"]:
        itype = item["type"]
        if itype == "text":
            p = _para(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            _add_bold(p, item["text"])
        elif itype == "image":
            _add_image(doc, item["path"])

    # 3. Висновок
    conc = data.get("conclusion")
    if conc:
        p = _para(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        _add_bold(p, f"**{conc[:9]}**{conc[9:]}")

    os.makedirs(os.path.dirname(out_docx), exist_ok=True)
    doc.save(out_docx)
    print(f"Збережено звіт: {out_docx}")


if __name__ == "__main__":
    build_docx(REPORT, OUT_DOCX)
