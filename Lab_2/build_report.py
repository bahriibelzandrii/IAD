# -*- coding: utf-8 -*-
"""Генератор офіційного звіту DOCX для Лабораторної роботи №2 (ІАД).

Вимоги:
- Шрифт: Times New Roman 14 pt, міжрядковий інтервал 1.15, поля 1.5 см.
- Графіки: 14.5 см, центровані.
- Титульний блок у родовому відмінку: Багрія-Белза Андрія, КН-2327б.
- Лаконічний зміст: чітка мета, опис кроків без води, всі 4 графіки та числові метрики, висновок від першої особи.
"""
import os
import re
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BASE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(BASE, "figures")
OUT_DIR = os.path.join(BASE, "report")
OUT_DOCX = os.path.join(OUT_DIR, "ІАД_КН-2327Б_Багрій-Белз_ЛР-2.docx")

BODY_FONT = "Times New Roman"
CODE_FONT = "Cascadia Mono"
BODY_SIZE = Pt(14)
CODE_SIZE = Pt(9.5)
FIG_WIDTH = Cm(14.5)
MARGIN_CM = 1.5


def f(name):
    return os.path.join(FIG, name)


META = {
    "title_word": "ЗВІТ",
    "lab_number": 2,
    "topic": "Регресійний аналіз даних. Кореляція",
    "course_name": "Інтелектуальний аналіз даних",
    "group": "КН-2327б",
    "author": "Багрій-Белз Андрій",
    "objective": "Дослідити лінійні взаємозв'язки між числовими змінними засобами Python, навчитися обчислювати коефіцієнт кореляції Пірсона, будувати парну лінійну регресійну модель, оцінювати її адекватність за метриками R² та MSE/RMSE, а також здійснювати статистичну діагностику залишків моделі на нормальність та гомоскедастичність."
}

BODY = [
    {
        "type": "task",
        "number": 1,
        "text": "Дослідження взаємозв'язку між довжиною тіла та вагою анаконд (основний набір даних anaconda.dat)."
    },
    {
        "type": "text",
        "text": (
            "Набір даних містить 56 спостережень дорослих анаконд (по 28 особин кожної статі). "
            "Категоріальну змінну статі закодовано бінарно: M = 0 (самці), F = 1 (самиці). "
            "Обчислено кореляційну матрицю Пірсона для довжини тіла (X, см), маси (Y, кг) та статі."
        )
    },
    {
        "type": "result",
        "label": "Матриця кореляції та параметри лінійної регресії anaconda.dat:",
        "images": []
    },
    {
        "type": "text",
        "text": (
            "Між довжиною тіла X та масою тіла Y виявлено надзвичайно сильний прямий лінійний зв'язок: "
            "**r = 0.9614** (p-value = **6.18e-32**, що свідчить про абсолютну статистичну значущість зв'язку). "
            "Кореляція між довжиною та закодованою статтю становить r = 0.7661, а між масою та статтю — r = 0.7299 (самиці помітно більші за самців).\n"
            "Побудовано модель парної лінійної регресії за методом найменших квадратів:"
        )
    },
    {
        "type": "formula",
        "text": "*ŷ* = 0.2530 · *X* − 50.7306"
    },
    {
        "type": "text",
        "text": (
            "де кутовий коефіцієнт a = 0.2530, вільний член b = -50.7306. "
            "Коефіцієнт детермінації становить **R² = 0.9243**, тобто побудована модель пояснює 92.43% дисперсії маси анаконди. "
            "Середньоквадратична помилка становить **MSE = 31.9250** (RMSE = **5.6502 кг**)."
        )
    },
    {
        "type": "result",
        "label": "Графік емпіричних даних та підігнаної прямої регресії:",
        "images": [f("reg_plot.png")]
    },
    {
        "type": "result",
        "label": "Діагностика залишків моделі (лінійність та гомоскедастичність):",
        "images": [f("residuals_plot.png")]
    },
    {
        "type": "text",
        "text": (
            "Середнє арифметичне значення залишків регресії становить практично нуль (-8.44e-15 кг), що підтверджує незміщеність оцінки МНК. "
            "Проте графік залишків Residuals vs Predictor X демонструє явну U-подібну форму та розширення дисперсії при зростанні довжини: "
            "для малих і великих значень довжини залишки переважно додатні, а для середніх — від'ємні. "
            "Це свідчить про наявність гетероскедастичності та нелінійного (наближеного до степеневого/кубічного) характеру залежності маси від довжини в біологічних об'єктах."
        )
    },
    {
        "type": "result",
        "label": "Гістограма розподілу залишків та оцінка нормальності:",
        "images": [f("residuals_hist.png")]
    },
    {
        "type": "text",
        "text": (
            "Оцінка нормальності розподілу залишків за критерієм Шапіро-Уїлка дала статистику W = 0.9176 (p-value = **9.67e-04** < 0.05). "
            "Це дозволяє відхилити нульову гіпотезу про нормальний розподіл залишків на користь правосторонньої асиметрії, "
            "що додатково аргументує доцільність використання нелінійної або степеневої алометричної моделі (наприклад, *Y* = *a* · *X*ᵇ) у практичних задачах біометрії."
        )
    },
    {
        "type": "task",
        "number": 2,
        "text": "Дослідження залежності між площею квартири та її ціною на власному числовому наборі даних flats.csv."
    },
    {
        "type": "text",
        "text": (
            "Для аналізу використано набір даних вторинного ринку нерухомості flats.csv (684 спостереження після очищення некоректних даних та викидів). "
            "Досліджено залежність між загальною площею квартири (незалежна змінна X, м²) та її вартістю (залежна змінна Y, грн)."
        )
    },
    {
        "type": "result",
        "label": "Парна регресія та кореляційний аналіз для ринку житла:",
        "images": [f("custom_regression.png")]
    },
    {
        "type": "text",
        "text": (
            "Коефіцієнт кореляції Пірсона між площею та ціною становить **r = 0.6729** при p-value = **2.41e-91**, "
            "що свідчить про помітний стійкий позитивний зв'язок зі стовідсотковою статистичною значущістю.\n"
            "Отримана лінійна модель має вигляд:"
        )
    },
    {
        "type": "formula",
        "text": "*Ціна* = 25180.99 · *Площа* − 408119.63"
    },
    {
        "type": "text",
        "text": (
            "де кутовий коефіцієнт a = 25180.99 грн/м² інтерпретується як гранична середня ціна одного додаткового квадратного метра житла. "
            "Коефіцієнт детермінації становить **R² = 0.4527**, RMSE = **801 528.39 грн**. "
            "Це демонструє, що загальна площа пояснює близько 45.3% варіації вартості квартир, тоді як решта 54.7% припадає на розташування (район/місто), поверх, клас забудови та стан ремонту."
        )
    }
]

CONCLUSION = (
    "На цій лабораторній роботі я засвоїв методи побудови та діагностики моделей парної лінійної регресії "
    "і кореляційного аналізу в Python з використанням бібліотек pandas, scipy.stats, scikit-learn, matplotlib та seaborn. "
    "Я навчився оцінювати якість регресії за коефіцієнтом детермінації R², середньоквадратичною помилкою MSE/RMSE "
    "та проводити діагностику ключових передумов МНК (лінійності, гомоскедастичності та нормальності розподілу залишків). "
    "В ході роботи встановлено, що для біометричних даних анаконд лінійна модель має високий показник R² = 0.9243, проте залишки свідчать про перевагу алометричної степеневої залежності, "
    "а на ринку житла площа виступає ключовим, але не єдиним ціноутворюючим фактором (R² = 0.4527)."
)

REPORT = {"meta": META, "body": BODY, "conclusion": CONCLUSION}


def _set_font(run, family, size, bold=False, italic=False):
    run.font.name = family
    run.font.size = size
    run.font.bold = bold
    run.font.italic = italic
    rPr = run._element.get_or_add_rPr()
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), family)
    rFonts.set(qn("w:hAnsi"), family)
    rFonts.set(qn("w:cs"), family)
    rFonts.set(qn("w:eastAsia"), family)
    rPr.append(rFonts)


def _para(doc, text="", *, align=None, font=BODY_FONT, size=BODY_SIZE,
          bold=False, line_spacing=1.15, keep_with_next=False, keep_together=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = keep_with_next
    p.paragraph_format.keep_together = keep_together
    if align is not None:
        p.alignment = align
    if text:
        _set_font(p.add_run(text), font, size, bold=bold)
    return p


LINK_RE = re.compile(r"\[([^\]]+)\]\((https?://[^)\s]+)\)")


TOKEN_RE = re.compile(r"(\*\*\*[^*]+\*\*\*|\*\*[^*]+\*\*|\*[^*]+\*)")


def _add_spans(p, text, font=BODY_FONT, size=BODY_SIZE):
    pos = 0
    for m in TOKEN_RE.finditer(text):
        if m.start() > pos:
            _set_font(p.add_run(text[pos:m.start()]), font, size, bold=False, italic=False)
        tok = m.group(1)
        if tok.startswith("***"):
            _set_font(p.add_run(tok[3:-3]), font, size, bold=True, italic=True)
        elif tok.startswith("**"):
            _set_font(p.add_run(tok[2:-2]), font, size, bold=True, italic=False)
        elif tok.startswith("*"):
            _set_font(p.add_run(tok[1:-1]), font, size, bold=False, italic=True)
        pos = m.end()
    if pos < len(text):
        _set_font(p.add_run(text[pos:]), font, size, bold=False, italic=False)

def _add_hyperlink(p, url, text, font=BODY_FONT, size=BODY_SIZE):
    part = p.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "004B87")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(color)
    rPr.append(u)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    hyperlink.append(r)
    p._p.append(hyperlink)


def _add_rich(p, text, font=BODY_FONT, size=BODY_SIZE):
    pos = 0
    for m in LINK_RE.finditer(text):
        if m.start() > pos:
            _add_spans(p, text[pos:m.start()], font, size)
        _add_hyperlink(p, m.group(2), m.group(1), font, size)
        pos = m.end()
    if pos < len(text):
        _add_spans(p, text[pos:], font, size)


def _add_image(doc, path):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = False
    p.paragraph_format.keep_together = True
    if os.path.exists(path):
        p.add_run().add_picture(path, width=FIG_WIDTH)
    else:
        print(f"  warning: image not found: {path}")


def _genitive(name):
    # Прізвище в родовому відмінку: Багрія-Белза Андрія
    return "Багрія-Белза Андрія"


def _title_block(doc, meta):
    c = WD_ALIGN_PARAGRAPH.CENTER
    lines = [
        meta.get("title_word", "ЗВІТ"),
        f"про виконання лабораторної роботи №{meta['lab_number']}",
        f"«{meta['topic']}»",
        "з дисципліни",
        f"«{meta['course_name']}»",
        f"Студента групи {meta['group']}",
        _genitive(meta["author"]),
    ]
    for i, line in enumerate(lines):
        _para(doc, line, align=c, bold=(i == len(lines) - 1))


def build_docx(data, out_docx):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(MARGIN_CM)
    sec.top_margin = sec.bottom_margin = Cm(MARGIN_CM)
    sec.header_distance = sec.footer_distance = Cm(1.25)
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(0)

    _title_block(doc, data["meta"])

    if data["meta"].get("objective"):
        p = _para(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY, keep_with_next=True)
        _set_font(p.add_run("Мета роботи: "), BODY_FONT, BODY_SIZE, bold=True)
        _add_rich(p, data["meta"]["objective"])

    J, L = WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.LEFT
    list_counter = 0
    for block in data["body"]:
        if block.get("page_break_before"):
            doc.add_page_break()
        btype = block["type"]
        if btype == "text":
            for line in block["text"].split("\n"):
                line = line.strip()
                if not line:
                    continue
                kwn = any(line.startswith(prefix) for prefix in ("**1.", "**2.", "**3.", "**4."))
                p = _para(doc, align=J, keep_with_next=kwn)
                _add_rich(p, line)
        elif btype == "task":
            list_counter = 0
            p = _para(doc, align=J, keep_with_next=True)
            _set_font(p.add_run(f"Завдання {block['number']}: "), BODY_FONT, BODY_SIZE, bold=True)
            _add_rich(p, block.get("text", ""))
        elif btype == "formula":
            p = _para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, keep_with_next=True)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            _add_rich(p, block["text"])
        elif btype == "section":
            _para(doc, block["text"], align=J, bold=True, keep_with_next=True)
        elif btype == "code":
            list_counter += 1
            _para(doc, f"{list_counter}. {block['label']}", align=L, bold=True, keep_with_next=True)
            for line in block["text"].splitlines():
                _para(doc, line, align=L, font=CODE_FONT, size=CODE_SIZE,
                      line_spacing=Pt(9.5))
        elif btype == "result":
            lbl = f"{list_counter + 1}. {block['label']}" if list_counter > 0 else block["label"]
            _para(doc, lbl, align=L, bold=True, keep_with_next=True)
            for img in block["images"]:
                _add_image(doc, img)
        else:
            raise ValueError(f"unknown block type: {btype!r}")

    p = _para(doc, align=J)
    _set_font(p.add_run("Висновок: "), BODY_FONT, BODY_SIZE, bold=True)
    _add_rich(p, data["conclusion"])

    Path(out_docx).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_docx)
    print(f"wrote {out_docx}")


if __name__ == "__main__":
    build_docx(REPORT, OUT_DOCX)
