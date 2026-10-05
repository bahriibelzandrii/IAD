# -*- coding: utf-8 -*-
"""
Build the IAD Lab 3 report (.docx) matching the Zvit house style.
"""
import os
import re
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BASE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(BASE, "figures")
OUT_DIR = os.path.join(BASE, "report")
OUT_DOCX = os.path.join(OUT_DIR, "ІАД_КН-2327Б_Багрій-Белз_ЛР-3.docx")

BODY_FONT = "Times New Roman"
BODY_SIZE = Pt(14)
FIG_WIDTH = Cm(15.0)
MARGIN_CM = 1.5


def f(name):
    return os.path.join(FIG, name)


# ---------------------------------------------------------------------------
# Content definition
# ---------------------------------------------------------------------------
META = {
    "course_name": "Інтелектуальний аналіз даних",
    "course_code": "ІАД",
    "lab_number": 3,
    "topic": "Дисперсійний аналіз даних (ANOVA)",
    "author": "Багрій-Белз Андрій",
    "group": "КН-2327Б",
    "objective": "Дослідити відмінності середніх значень ознак під впливом одного та кількох якісних факторів за допомогою однофакторного і двофакторного дисперсійного аналізу в Python."
}

BODY = [
    {
        "type": "task",
        "number": 1,
        "text": "Дослідження зв'язку тембру голосу та зросту співаків (Voice.txt, One-Way ANOVA)"
    },
    {
        "type": "text",
        "text": (
            "Для перевірки залежності зросту співаків від типу голосу було побудовано лінійну модель "
            "та проведено однофакторний дисперсійний аналіз (One-Way ANOVA). Для повної вибірки співаків отримано статистику "
            "**F = 54.31** та досягнутий рівень значущості **p = 1.33 × 10⁻²²** (p < 0.05). "
            "З метою з'ясування природи цього ефекту проведено роздільний дисперсійний аналіз у межах кожної статі. "
            "Для жіночої підгрупи (Soprano та Alto) отримано **F = 1.44**, **p = 0.235** (p > 0.05). "
            "Для чоловічої підгрупи (Tenor та Bass) отримано **F = 0.48**, **p = 0.492** (p > 0.05). "
            "Оскільки всередині кожної статевої групи відмінності зросту за тембром голосу є статистично незначущими, "
            "виявлена на повній вибірці залежність є виключно артефактом статевого диморфізму (жінки в середньому нижчі за чоловіків)."
        )
    },
    {
        "type": "result",
        "images": [f("voice_boxplot.png")]
    },
    {
        "type": "task",
        "number": 2,
        "text": "Дослідження кількості помилок щурів у лабіринті (rat.txt, Two-Way ANOVA)"
    },
    {
        "type": "text",
        "text": (
            "Для оцінки одночасного впливу умов середовища (ENVIRNMNT) та генетичної лінії (STRAIN) щурів на кількість "
            "помилок у лабіринті (ERRORS) виконано двофакторний дисперсійний аналіз типу II (Type II Two-Way ANOVA). "
            "Фактор середовища виявився статистично значущим: **F = 5.82**, **p = 0.0267** (p < 0.05). "
            "Фактор генетичної лінії також є статистично значущим: **F = 4.16**, **p = 0.0326** (p < 0.05). "
            "Взаємодія факторів (ENVIRNMNT:STRAIN) є повністю незначущою: **F = 0.0084**, **p = 0.9916** (p > 0.05). "
            "Це свідчить про суто адитивний вплив середовища та спадковості: умови обмеження підвищують кількість помилок "
            "однаковою мірою для кожної генетичної лінії, що підтверджується паралельністю ліній на графіку взаємодії."
        )
    },
    {
        "type": "result",
        "images": [f("rat_interaction.png"), f("rat_boxplot.png")]
    },
    {
        "type": "task",
        "number": 3,
        "text": "Дисперсійний аналіз власного набору даних (Tips dataset, One-Way ANOVA)"
    },
    {
        "type": "text",
        "text": (
            "Досліджено вплив дня тижня (day: Thur, Fri, Sat, Sun) на загальну суму чека (total_bill) відвідувачів ресторану. "
            "Попередню перевірку гомогенності дисперсій між групами здійснено за тестом Левена: **W = 0.665**, **p = 0.574** (p > 0.05), "
            "що підтверджує коректність застосування дисперсійного аналізу. "
            "За результатами One-Way ANOVA встановлено статистично значущий вплив дня тижня на величину чека: "
            "**F = 2.77**, **p = 0.0425** (p < 0.05). У вихідні дні (субота та неділя) спостерігається вищий середній рівень витрат клієнтів."
        )
    },
    {
        "type": "result",
        "images": [f("custom_anova.png")]
    }
]

CONCLUSION = (
    "На цій лабораторній роботі я опанував практичне застосування однофакторного та двофакторного дисперсійного аналізу (ANOVA) "
    "в екосистемі Python за допомогою бібліотек statsmodels та scipy. "
    "Я навчився оцінювати вплив якісних факторів та їхніх взаємодій на кількісні показники, перевіряти передумови гомогенності дисперсій, "
    "а також коректно виявляти приховані фактори у дослідженнях, що запобігає хибним висновкам через побічні артефакти вибірки."
)

REPORT = {"meta": META, "body": BODY, "conclusion": CONCLUSION}


# ---------------------------------------------------------------------------
# Formatting helpers
# ---------------------------------------------------------------------------
def _set_font(run, family, size, bold=False):
    run.font.name = family
    run.font.size = size
    run.bold = bold
    rPr = run._r.get_or_add_rPr()
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), family)
    rFonts.set(qn("w:hAnsi"), family)
    rFonts.set(qn("w:eastAsia"), family)
    rFonts.set(qn("w:cs"), family)


def _para(doc, text="", *, align=None, font=BODY_FONT, size=BODY_SIZE,
          bold=False, line_spacing=1.15, keep_with_next=False, keep_together=True):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if align is not None:
        pf.alignment = align
    pf.line_spacing = line_spacing
    pf.space_before = Pt(0)
    pf.space_after = Pt(4)
    pf.keep_with_next = keep_with_next
    pf.keep_together = keep_together
    if text:
        run = p.add_run(text)
        _set_font(run, font, size, bold=bold)
    return p


def _add_bold(p, text, font=BODY_FONT, size=BODY_SIZE):
    parts = text.split("**")
    for i, seg in enumerate(parts):
        if not seg:
            continue
        if i % 2 == 1:
            _set_font(p.add_run(seg), font, size, bold=True)
        else:
            _set_font(p.add_run(seg), font, size, bold=False)


def _add_image(doc, path):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_together = True
    if os.path.exists(path):
        run = p.add_run()
        run.add_picture(path, width=FIG_WIDTH)
    else:
        print(f"  warning: image not found: {path}")


def _genitive(name):
    # 'Багрій-Белз Андрій' -> 'Багрія-Белза Андрія'
    parts = name.split()
    if len(parts) < 2:
        return name
    sur, first = parts[0], parts[1]
    # For compound surname 'Багрій-Белз'
    sur_sub = []
    for sub in sur.split("-"):
        if sub.endswith(("й", "ї")):
            sur_sub.append(sub[:-1] + "я")
        elif not sub.endswith(("а", "я")):
            sur_sub.append(sub + "а")
        else:
            sur_sub.append(sub)
    sur_gen = "-".join(sur_sub)
    
    if first.endswith(("й", "ї")):
        first_gen = first[:-1] + "я"
    elif not first.endswith(("а", "я")):
        first_gen = first + "а"
    else:
        first_gen = first
    return f"{sur_gen} {first_gen}"


def _title_block(doc, meta):
    c = WD_ALIGN_PARAGRAPH.CENTER
    lines = [
        "ЗВІТ",
        f"про виконання лабораторної роботи №{meta['lab_number']}",
        f"«{meta['topic']}»",
        "з дисципліни",
        f"«{meta['course_name']}»",
        f"Студента групи {meta['group']}",
        _genitive(meta["author"]),
    ]
    for i, line in enumerate(lines):
        _para(doc, line, align=c, bold=(i == len(lines) - 1))


def build_docx(data, out_docx):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(MARGIN_CM)
    sec.top_margin = sec.bottom_margin = Cm(MARGIN_CM)
    sec.header_distance = sec.footer_distance = Cm(1.25)
    
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(0)

    _title_block(doc, data["meta"])

    if data["meta"].get("objective"):
        p = _para(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY, keep_with_next=True)
        _set_font(p.add_run("Мета роботи: "), BODY_FONT, BODY_SIZE, bold=True)
        _add_bold(p, data["meta"]["objective"])

    J = WD_ALIGN_PARAGRAPH.JUSTIFY
    for block in data["body"]:
        btype = block["type"]
        if btype == "task":
            p = _para(doc, align=J, keep_with_next=True)
            _set_font(p.add_run(f"Завдання {block['number']}: "), BODY_FONT, BODY_SIZE, bold=True)
            _add_bold(p, block.get("text", ""))
        elif btype == "text":
            p = _para(doc, align=J, keep_with_next=True)
            _add_bold(p, block["text"])
        elif btype == "result":
            for img in block["images"]:
                _add_image(doc, img)

    p = _para(doc, align=J)
    _set_font(p.add_run("Висновок: "), BODY_FONT, BODY_SIZE, bold=True)
    _add_bold(p, data["conclusion"])

    Path(out_docx).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_docx)
    print(f"Збережено звіт: {out_docx}")


if __name__ == "__main__":
    build_docx(REPORT, OUT_DOCX)
