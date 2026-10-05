# -*- coding: utf-8 -*-
"""
Build the IAD Lab 6 report (.docx) matching the canonical house style.
"""
import os
import re
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BASE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(BASE, "figures")
OUT_DIR = os.path.join(BASE, "report")
OUT_DOCX = os.path.join(OUT_DIR, "ІАД_КН-2327Б_Багрій-Белз_ЛР-6.docx")

BODY_FONT = "Times New Roman"
BODY_SIZE = Pt(14)
FIG_WIDTH = Cm(14.5)
MARGIN_CM = 1.5


def f(name):
    return os.path.join(FIG, name)


META = {
    "course_name": "Інтелектуальний аналіз даних",
    "course_code": "ІАД",
    "lab_number": 6,
    "topic": "Використання самоорганізуючихся карт Кохонена (SOM)",
    "author": "Багрій-Белз Андрій",
    "group": "КН-2327Б",
    "objective": (
        "Дослідити архітектуру та практичне застосування самоорганізуючихся карт Кохонена (Self-Organizing Maps, SOM) "
        "для розв'язання задач неконтрольованого аналізу даних: топологічної кластеризації, проєкції багатовимірного простору ознак "
        "на двовимірну решітку нейронів, аналізу щільності розподілу та виявлення нелінійних взаємозв'язків між факторами."
    )
}

TRAIN_DESC = (
    "Для дослідження використано набір даних Boston Housing (506 спостережень, 7 числових показників: "
    "INDUS, DIS, NOX, LSTAT, AGE, RAD, B), попередньо стандартизованих за допомогою StandardScaler. "
    "Ініціалізовано двовимірну решітку нейронів розміром **9 x 6** (54 нейрони) з параметрами sigma=1.0, learning_rate=0.5. "
    "Модель навчено протягом 100 ітерацій пакетного навчання (batch training). "
    "Динаміка похибки квантування (Quantization Error) демонструє стійке монотонне збігання з 1.8212 до **1.7809** "
    "при збереженні топологічної структури простору."
)

UMATRIX_DESC = (
    "Побудовано U-Matrix (Unified Distance Matrix), що візуалізує евклідові відстані між ваговими векторами сусідніх нейронів "
    "у латентному просторі. Світлі ділянки з малими міжнейронними відстанями відповідають ядрам високої щільності та кластерній "
    "однорідності об'єктів, тоді як темні контрастні смуги виступають топологічними бар'єрами (межами) між різними типами житлових районів міста.\n"
    "Матриця активацій нейронів (Activation Frequency / Counts) відображає кількість спостережень, для яких кожен нейрон став найкращим переможцем (BMU). "
    "Розподіл частот підтверджує рівномірне покриття навчальної вибірки решіткою: спостерігаються виражені вузли фокусування спостережень "
    "(до 23 об'єктів на нейрон) та відсутність надмірних «мертвих» зон, що свідчить про якісну адаптацію топологічної сітки."
)

PLANES_DESC = (
    "Показано проєкцію всіх 506 спостережень на карту Кохонена поверх U-Matrix із колірним кодуванням за показником **LSTAT** "
    "(% населення з низьким соціально-економічним статусом). Спостерігається чітке просторове розшарування: об'єкти з високим LSTAT "
    "зосереджені у верхньому та лівому секторах решітки, тоді як престижні райони з низьким рівнем LSTAT локалізуються в протилежній "
    "частині карти, розділяючись природними бар'єрами U-Matrix.\n"
    "Аналіз компонентних площин (Component Planes) дозволяє зіставити просторові патерни ваг нейронів за окремими змінними. "
    "Виявлено виражену синхронну поведінку (сильну пряму кореляцію) факторів NOX (забруднення оксидами азоту), INDUS (частка промислових зон), "
    "AGE (вік забудови) та LSTAT — їхні максимуми припадають на верхні та ліві нейрони. Натомість показник DIS (відстань до центрів зайнятості) "
    "демонструє дзеркальний розподіл значень, що підтверджує нелінійну обернену залежність між урбанізацією, екологічним станом та віддаленістю від центру."
)

CONCLUSION = (
    "На цій лабораторній роботі я опанував метод самоорганізуючихся карт Кохонена (SOM) для розвідувального аналізу даних у Python. "
    "Успішно навчив нейромережу розмірністю 9 x 6 на стандартизованих ознаках вибірки Boston Housing, зафіксувавши зниження "
    "похибки квантування до 1.7809. Завдяки спільній візуалізації U-Matrix, матриці активацій та компонентних площин я виявив "
    "топологічні кластери спостережень і наочно підтвердив наявність взаємозв'язків між промисловим забрудненням, соціальним статусом "
    "та просторовою віддаленістю житлових масивів."
)

BODY = [
    {
        "type": "task",
        "number": 1,
        "text": "Побудова та навчання мережі SOM, динаміка збіжності"
    },
    {
        "type": "text",
        "text": TRAIN_DESC
    },
    {
        "type": "result",
        "images": [f("som_training_loss.png")]
    },
    {
        "type": "task",
        "number": 2,
        "text": "Візуалізація карти відстаней (U-Matrix) та частоти активацій"
    },
    {
        "type": "text",
        "text": UMATRIX_DESC
    },
    {
        "type": "result",
        "images": [f("som_umatrix.png"), f("som_counts.png")]
    },
    {
        "type": "task",
        "number": 3,
        "text": "Проєкція об'єктів (BMU) та компонентні карти"
    },
    {
        "type": "text",
        "text": PLANES_DESC
    },
    {
        "type": "result",
        "images": [f("som_bmu_mapping.png"), f("som_component_planes.png")]
    }
]

REPORT = {"meta": META, "body": BODY, "conclusion": CONCLUSION}


def _set_font(run, family, size, bold=False):
    run.font.name = family
    run.font.size = size
    run.bold = bold
    rPr = run._r.get_or_add_rPr()
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), family)
    rFonts.set(qn("w:hAnsi"), family)
    rFonts.set(qn("w:eastAsia"), family)
    rFonts.set(qn("w:cs"), family)


def _para(doc, text="", *, align=None, font=BODY_FONT, size=BODY_SIZE,
          bold=False, line_spacing=1.15, keep_with_next=False, keep_together=True):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if align is not None:
        pf.alignment = align
    pf.line_spacing = line_spacing
    pf.space_before = Pt(0)
    pf.space_after = Pt(4)
    pf.keep_with_next = keep_with_next
    pf.keep_together = keep_together
    if text:
        run = p.add_run(text)
        _set_font(run, font, size, bold=bold)
    return p


def _add_bold(p, text, font=BODY_FONT, size=BODY_SIZE):
    parts = text.split("**")
    for i, seg in enumerate(parts):
        if not seg:
            continue
        if i % 2 == 1:
            _set_font(p.add_run(seg), font, size, bold=True)
        else:
            _set_font(p.add_run(seg), font, size, bold=False)


def _add_image(doc, path):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_together = True
    if os.path.exists(path):
        run = p.add_run()
        run.add_picture(path, width=FIG_WIDTH)
    else:
        print(f"  warning: image not found: {path}")


def _genitive(name):
    parts = name.split()
    if len(parts) < 2:
        return name
    sur, first = parts[0], parts[1]
    sur_sub = []
    for sub in sur.split("-"):
        if sub.endswith(("й", "ї")):
            sur_sub.append(sub[:-1] + "я")
        elif not sub.endswith(("а", "я")):
            sur_sub.append(sub + "а")
        else:
            sur_sub.append(sub)
    sur_gen = "-".join(sur_sub)

    if first.endswith(("й", "ї")):
        first_gen = first[:-1] + "я"
    elif not first.endswith(("а", "я")):
        first_gen = first + "а"
    else:
        first_gen = first
    return f"{sur_gen} {first_gen}"


def _title_block(doc, meta):
    c = WD_ALIGN_PARAGRAPH.CENTER
    lines = [
        "ЗВІТ",
        f"про виконання лабораторної роботи №{meta['lab_number']}",
        f"«{meta['topic']}»",
        "з дисципліни",
        f"«{meta['course_name']}»",
        f"Студента групи {meta['group']}",
        _genitive(meta["author"]),
    ]
    for i, line in enumerate(lines):
        _para(doc, line, align=c, bold=(i == len(lines) - 1))


def build_docx(data, out_docx):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(MARGIN_CM)
    sec.top_margin = sec.bottom_margin = Cm(MARGIN_CM)
    sec.header_distance = sec.footer_distance = Cm(1.25)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(0)

    _title_block(doc, data["meta"])

    if data["meta"].get("objective"):
        p = _para(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY, keep_with_next=True)
        _set_font(p.add_run("Мета роботи: "), BODY_FONT, BODY_SIZE, bold=True)
        _add_bold(p, data["meta"]["objective"])

    J = WD_ALIGN_PARAGRAPH.JUSTIFY
    for block in data["body"]:
        btype = block["type"]
        if btype == "task":
            p = _para(doc, align=J, keep_with_next=True)
            _set_font(p.add_run(f"Завдання {block['number']}: "), BODY_FONT, BODY_SIZE, bold=True)
            _add_bold(p, block.get("text", ""))
        elif btype == "text":
            for line in block["text"].split("\n"):
                line = line.strip()
                if not line:
                    continue
                p = _para(doc, align=J, keep_with_next=True)
                _add_bold(p, line)
        elif btype == "result":
            for img in block["images"]:
                _add_image(doc, img)

    p = _para(doc, align=J)
    _set_font(p.add_run("Висновок: "), BODY_FONT, BODY_SIZE, bold=True)
    _add_bold(p, data["conclusion"])

    Path(out_docx).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_docx)
    print(f"Збережено звіт: {out_docx}")


if __name__ == "__main__":
    build_docx(REPORT, OUT_DOCX)
