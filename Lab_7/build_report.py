# -*- coding: utf-8 -*-
"""Скрипт генерації Lab_7/build_report.py відповідно до канонічного корпоративного стилю звіту (report-maker)."""

import json
import os
import re

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(BASE, "figures")
OUT_DIR = os.path.join(BASE, "report")
OUT_DOCX = os.path.join(OUT_DIR, "ІАД_КН-2327Б_Багрій-Белз_ЛР-7.docx")

BODY_FONT = "Times New Roman"
CODE_FONT = "Cascadia Mono"
BODY_SIZE = Pt(14)
CODE_SIZE = Pt(9.5)
FIG_WIDTH = Cm(15.0)
MARGIN_CM = 1.5


def f(name):
    return os.path.join(FIG, name)


META = {
    "course_code": "ІАД",
    "course_name": "Інтелектуальний аналіз даних",
    "lab_number": "7",
    "topic": "Прогнозування з використанням часових рядів",
    "author": "Багрія-Белза Андрія",
    "group": "КН-2327Б",
    "title_word": "ЗВІТ",
    "title_style": "standard",
}

PURPOSE = (
    "Дослідити структуру часових рядів, опанувати методи згладжування простим ковзним середнім (SMA) "
    "та адитивної сезонної декомпозиції, реалізувати хронологічний розподіл вибірки та порівняти точність "
    "базового наївного прогнозу (NaiveForecaster) і моделі експоненційного згладжування Гольта-Вінтерса "
    "(ExponentialSmoothing) за показниками MAE, RMSE та MAPE."
)

SMA_DESC = (
    "Виконано згладжування хронологічного ряду тривалості життя 42 монархів Англії (kings.txt) "
    "за допомогою простого ковзного середнього (Simple Moving Average, n=9) із центрованим вікном. "
    "Фактичні значення мають значний розмах (від 13 до 86 років). Застосування центрованого ковзного вікна "
    "ефективно фільтрує випадкові високочастотні стрибки та чітко візуалізує стабільний висхідний тренд "
    "тривалості життя королів у пізніші історичні періоди."
)

DECOMP_DESC = (
    "Здійснено класичну адитивну декомпозицію часового ряду народжуваності (babyboom.txt, 168 спостережень) "
    "із періодом s=12 на 4 складові: спостережуваний ряд (Observed), тренд (Trend), сезонність (Seasonal) "
    "та випадковий шум (Residual). Графік підтверджує наявність плавної нелінійної трендової траєкторії, "
    "строгої повторюваної річної сезонності з амплітудою коливань від -2.08 до +1.46 од., "
    "а також гомоскедастичних залишкових похибок, близьких до білого шуму."
)

SPLIT_DESC = (
    "Виконано завантаження щомісячного ряду міжнародних авіаперевезень Airline Passengers (144 місяці, 1949–1960 рр.) "
    "та здійснено хронологічний розподіл на тренувальну (Train, 108 перших місяців) і тестову (Test, 36 останніх місяців) "
    "вибірки за допомогою temporal_train_test_split. Такий підхід унеможливлює витік даних з майбутнього "
    "та забезпечує об'єктивну оцінку екстраполяційних можливостей прогнозних алгоритмів."
)

FORECAST_DESC = (
    "Згенеровано прогноз на горизонт fh=36 місяців за двома моделями: базовою NaiveForecaster (стратегія 'last') "
    "та адаптивною моделлю експоненційного згладжування Гольта-Вінтерса (адитивний тренд і адитивна сезонність s=12). "
    "Суміщений графік демонструє, що наївна модель фіксує лише константу останньої точки (y=336), ігноруючи ріст і сезонність, "
    "тоді як модель Гольта-Вінтерса точно повторює фази літніх максимумів і загальний висхідний нахил тренду."
)

METRICS_INTRO = (
    "Для кількісного зіставлення якості прогнозних підходів на тестовій вибірці (36 місяців) "
    "розраховано метрики MAE, RMSE та MAPE. Результати зведено у таблицю 1."
)

CONCLUSION = (
    "Під час виконання лабораторної роботи я опанував практичні методи аналізу та прогнозування часових рядів у середовищі Python. "
    "Дослідив структуру рядів за допомогою згладжування ковзним середнім (SMA) та адитивної декомпозиції, виділивши ортогональні компоненти тренду, сезонності та залишків. "
    "Провівши експериментальне прогнозування на тестовому горизонті 36 місяців, я переконався у перевазі моделі експоненційного згладжування Гольта-Вінтерса "
    "над наївним підходом: урахування сезонності та тренду дозволило знизити середню абсолютну відносну похибку (MAPE) з 19.89% до 5.11%."
)


def _set_font(run, family, size, bold=False, italic=False, color=None):
    run.font.name = family
    run.font.size = size
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), family)
    rFonts.set(qn("w:hAnsi"), family)
    rFonts.set(qn("w:cs"), family)


def _para(doc, text="", *, align=None, font=BODY_FONT, size=BODY_SIZE,
          bold=False, line_spacing=1.15, keep_with_next=False, keep_together=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    if align is not None:
        p.alignment = align
    if keep_with_next:
        p.paragraph_format.keep_with_next = True
    if keep_together:
        p.paragraph_format.keep_together = True
    if text:
        r = p.add_run(text)
        _set_font(r, font, size, bold=bold)
    return p


def _add_rich(p, text, font=BODY_FONT, size=BODY_SIZE):
    pattern = re.compile(r"(\*\*.*?\*\*)")
    tokens = pattern.split(text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            r = p.add_run(token[2:-2])
            _set_font(r, font, size, bold=True)
        else:
            r = p.add_run(token)
            _set_font(r, font, size, bold=False)


def _add_image(doc, path, width=FIG_WIDTH):
    if not os.path.exists(path):
        print(f"Помилка: файл зображення не знайдено: {path}")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_together = True
    r = p.add_run()
    r.add_picture(path, width=width)


def _set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, val in [("w:top", top), ("w:bottom", bottom), ("w:left", left), ("w:right", right)]:
        node = OxmlElement(m)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def _set_table_borders(table):
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement("w:tblBorders")
    for border_name in ["top", "left", "bottom", "right", "insideH"]:
        border = OxmlElement(f"w:{border_name}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "B0B0B0")
        tblBorders.append(border)
    insideV = OxmlElement("w:insideV")
    insideV.set(qn("w:val"), "none")
    tblBorders.append(insideV)
    tblPr.append(tblBorders)


def _add_metrics_table(doc, report_json_path):
    with open(report_json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(2)
    _add_rich(p_cap, "**Таблиця 1.** Порівняння метрик точності прогнозних моделей (горизонт fh = 36 місяців)")

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(table)

    headers = ["Модель прогнозування", "MAE", "RMSE", "MAPE (%)"]
    widths = [Cm(7.2), Cm(2.4), Cm(2.4), Cm(2.6)]

    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].width = widths[i]
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        for r in p.runs:
            _set_font(r, BODY_FONT, Pt(11), bold=True)

        # Тло заголовка таблиці
        tcPr = hdr_cells[i]._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "EAECEE")
        tcPr.append(shd)

    models_dict = data["models"]
    rows_data = [
        (
            models_dict["NaiveForecaster_last"]["name"],
            f"{models_dict['NaiveForecaster_last']['MAE']:.3f}",
            f"{models_dict['NaiveForecaster_last']['RMSE']:.3f}",
            f"{models_dict['NaiveForecaster_last']['MAPE_percent']:.2f}%",
        ),
        (
            models_dict["HoltWinters_ExponentialSmoothing"]["name"],
            f"{models_dict['HoltWinters_ExponentialSmoothing']['MAE']:.3f}",
            f"{models_dict['HoltWinters_ExponentialSmoothing']['RMSE']:.3f}",
            f"{models_dict['HoltWinters_ExponentialSmoothing']['MAPE_percent']:.2f}%",
        ),
    ]

    for m_name, mae, rmse, mape in rows_data:
        row_cells = table.add_row().cells
        vals = [m_name, mae, rmse, mape]
        for i, val in enumerate(vals):
            row_cells[i].text = val
            row_cells[i].width = widths[i]
            row_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_margins(row_cells[i], top=100, bottom=100, left=140, right=140)
            p = row_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                _set_font(r, BODY_FONT, Pt(11), bold=(i == 0))


def _title_block(doc, meta):
    c = WD_ALIGN_PARAGRAPH.CENTER
    _para(doc, meta["title_word"], align=c, bold=True)
    _para(doc, f"про виконання лабораторної роботи №{meta['lab_number']}", align=c)
    _para(doc, f"«{meta['topic']}»", align=c, bold=True)
    _para(doc, "з дисципліни", align=c)
    _para(doc, f"«{meta['course_name']}»", align=c)
    _para(doc, f"Студента групи {meta['group']}", align=c)
    _para(doc, meta["author"], align=c, bold=True)


def build_docx(out_docx):
    os.makedirs(os.path.dirname(out_docx), exist_ok=True)
    doc = Document()

    for sec in doc.sections:
        sec.top_margin = Cm(MARGIN_CM)
        sec.bottom_margin = Cm(MARGIN_CM)
        sec.left_margin = Cm(MARGIN_CM)
        sec.right_margin = Cm(MARGIN_CM)

    # 1. Титульний блок
    _title_block(doc, META)

    # 2. Мета роботи
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.line_spacing = 1.15
    p_meta.paragraph_format.space_before = Pt(6)
    p_meta.paragraph_format.space_after = Pt(6)
    p_meta.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _add_rich(p_meta, f"**Мета роботи:** {PURPOSE}")

    # 3. Завдання 1: Згладжування та декомпозиція
    p_t1 = doc.add_paragraph()
    p_t1.paragraph_format.line_spacing = 1.15
    p_t1.paragraph_format.space_before = Pt(6)
    p_t1.paragraph_format.space_after = Pt(4)
    p_t1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _add_rich(p_t1, f"**Завдання 1 (Згладжування та декомпозиція часових рядів).** {SMA_DESC}")

    _add_image(doc, f("kings_sma.png"), width=Cm(15.0))

    p_t1_d = doc.add_paragraph()
    p_t1_d.paragraph_format.line_spacing = 1.15
    p_t1_d.paragraph_format.space_before = Pt(4)
    p_t1_d.paragraph_format.space_after = Pt(4)
    p_t1_d.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _add_rich(p_t1_d, DECOMP_DESC)

    _add_image(doc, f("babyboom_decomposition.png"), width=Cm(15.0))

    # 4. Завдання 2: Прогнозування часових рядів
    p_t2 = doc.add_paragraph()
    p_t2.paragraph_format.line_spacing = 1.15
    p_t2.paragraph_format.space_before = Pt(6)
    p_t2.paragraph_format.space_after = Pt(4)
    p_t2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _add_rich(p_t2, f"**Завдання 2 (Прогнозування часових рядів).** {SPLIT_DESC}")

    _add_image(doc, f("split_series.png"), width=Cm(15.0))

    p_t2_f = doc.add_paragraph()
    p_t2_f.paragraph_format.line_spacing = 1.15
    p_t2_f.paragraph_format.space_before = Pt(4)
    p_t2_f.paragraph_format.space_after = Pt(4)
    p_t2_f.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _add_rich(p_t2_f, FORECAST_DESC)

    _add_image(doc, f("forecast_comparison.png"), width=Cm(15.0))

    # Оцінка точності та таблиця
    p_tbl_intro = doc.add_paragraph()
    p_tbl_intro.paragraph_format.line_spacing = 1.15
    p_tbl_intro.paragraph_format.space_before = Pt(4)
    p_tbl_intro.paragraph_format.space_after = Pt(4)
    p_tbl_intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _add_rich(p_tbl_intro, METRICS_INTRO)

    report_json_path = os.path.join(BASE, "report.json")
    _add_metrics_table(doc, report_json_path)

    # 5. Висновок
    p_concl = doc.add_paragraph()
    p_concl.paragraph_format.line_spacing = 1.15
    p_concl.paragraph_format.space_before = Pt(8)
    p_concl.paragraph_format.space_after = Pt(4)
    p_concl.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _add_rich(p_concl, f"**Висновок:** {CONCLUSION}")

    doc.save(out_docx)
    print(f"Згенеровано звіт: {out_docx}")


if __name__ == "__main__":
    build_docx(OUT_DOCX)
